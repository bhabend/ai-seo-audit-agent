"""A refused robots.txt or sitemap is a refusal, not an absent file.

The crawl path learned this first: a response the server refused is never
parsed or scored. The same site refuses robots.txt and the sitemap too, and
the run went on to say "the sitemap lists 0 addresses" and to treat the
robots file as absent. Neither is known. Both are claims about files nobody
was allowed to see.
"""

import json

import pytest

from conftest import BASE, crawl_site, html_page
from seo_audit.findings import build_findings
from seo_audit.issues import CRAWLER_EFFECT, is_blocked_status
from seo_audit.report import generate_short

# The same body the real edge network returned, as in test_blocked.py.
ACCESS_DENIED = (
    "<HTML><HEAD>\n<TITLE>Access Denied</TITLE>\n</HEAD><BODY>\n"
    "<H1>Access Denied</H1>\n \nYou don't have permission to access "
    "&#34;http&#58;&#47;&#47;example&#46;com&#47;&#34; on this server.<P>\n"
    "Reference&#32;&#35;18&#46;b20e0317&#46;1788946408&#46;2ff95345\n"
    "<P>https&#58;&#47;&#47;errors&#46;edgesuite&#46;net&#47;18&#46;"
    "b20e0317&#46;1788946408&#46;2ff95345</P>\n</BODY>\n</HTML>\n")

REFUSAL = {"status_code": 403, "text": ACCESS_DENIED,
           "headers": {"Content-Type": "text/html"}}


def refused_discovery(tmp_path, pages=None, **kwargs):
    """A site that serves pages but refuses robots.txt and every sitemap."""
    def extra(mock):
        for path in ("/robots.txt", "/sitemap.xml", "/sitemap_index.xml"):
            mock.get(BASE + path, **REFUSAL)
            mock.head(BASE + path, **REFUSAL)

    return crawl_site(tmp_path, pages or {BASE + "/": html_page()},
                      extra=extra, **kwargs)


# --- the refusal has a name of its own ---------------------------------------

def test_a_refusal_is_its_own_kind_of_finding():
    """"Nothing answered" and "we were refused" are different problems."""
    assert "request_refused" in CRAWLER_EFFECT
    assert "refused" in CRAWLER_EFFECT["request_refused"]
    # fetch_error goes back to meaning a transport failure and nothing else.
    assert "refused" not in CRAWLER_EFFECT["fetch_error"]
    assert "nothing back" in CRAWLER_EFFECT["fetch_error"]


def test_the_new_type_is_weighted_like_every_other_crawl_type():
    """Registered so the build passes, weighted so no score moves."""
    from seo_audit.scoring import (check_registry_coverage, load_weights,
                                   score_page)
    from seo_audit.issues import PAGE_ISSUE_SEVERITY

    weights = load_weights()
    entry = weights.crawl_issues["request_refused"]
    assert entry["bucket"] is None and entry["points"] == 0
    assert check_registry_coverage(weights, PAGE_ISSUE_SEVERITY,
                                   CRAWLER_EFFECT) == []
    assert score_page([], weights).score == 100


def test_a_refused_page_carries_the_new_type_not_fetch_error(tmp_path):
    def extra(mock):
        mock.get(BASE + "/locked", **REFUSAL)
        mock.head(BASE + "/locked", **REFUSAL)

    run = crawl_site(tmp_path, {BASE + "/": html_page(["/locked"])},
                     extra=extra)

    assert [row["url"] for row in run.issues_of("request_refused")] == \
        [BASE + "/locked"]
    assert run.issues_of("fetch_error") == []


# --- robots.txt --------------------------------------------------------------

def test_a_refused_robots_file_is_reported_as_refused(tmp_path):
    run = refused_discovery(tmp_path)
    refusals = run.issues_of("request_refused")
    robots = [row for row in refusals if row["url"].endswith("/robots.txt")]

    assert robots, refusals
    assert "403" in robots[0]["detail"]
    assert "refused" in robots[0]["detail"]
    assert run.summary["robots"]["refused"] is True
    assert run.summary["robots"]["status_code"] == 403


def test_a_missing_robots_file_is_still_merely_missing(tmp_path):
    """A 404 says the file is not there. That is a different fact."""
    run = crawl_site(tmp_path, {BASE + "/": html_page()})

    assert run.summary["robots"]["found"] is False
    assert run.summary["robots"]["refused"] is False
    assert [row for row in run.issues_of("request_refused")
            if row["url"].endswith("/robots.txt")] == []


def test_no_robots_finding_is_drawn_from_a_file_nobody_read(tmp_path):
    """The AI crawler rules finding reads robots.txt. There was no robots."""
    run = refused_discovery(tmp_path)

    assert "ai_crawler_rule" not in run.issue_types()
    assert "robots_blocked_linked" not in run.issue_types()
    assert "robots_blocked_in_sitemap" not in run.issue_types()
    assert run.summary["robots"]["disallow_rules"] == 0


# --- the sitemap -------------------------------------------------------------

def test_a_refused_sitemap_request_is_reported_as_refused(tmp_path):
    run = refused_discovery(tmp_path)
    sitemaps = [row for row in run.issues_of("request_refused")
                if "sitemap" in row["url"]]

    assert sitemaps, run.issues
    assert all("403" in row["detail"] for row in sitemaps)
    assert all("refused" in row["detail"] for row in sitemaps)
    refused = run.summary["sitemap"]["refused"]
    assert refused and all(entry["status"] == 403 for entry in refused)


def test_the_findings_carry_the_sitemap_refusal(tmp_path):
    run = refused_discovery(tmp_path)
    findings = build_findings(run.out_dir, compare_enabled=False)
    sitemap = findings["crawlability"]["sitemap"]

    assert sitemap["refused"], sitemap
    assert sitemap["requests"] >= len(sitemap["refused"])
    assert findings["crawlability"]["robots"]["refused"] is True
    assert findings["crawlability"]["robots"]["status_code"] == 403


def test_the_report_does_not_claim_an_empty_sitemap(tmp_path):
    """"The sitemap lists 0 addresses" is a claim about a refused file."""
    from docx import Document

    run = refused_discovery(tmp_path)
    findings = build_findings(run.out_dir, compare_enabled=False)
    result = generate_short(run.out_dir)
    document = Document(result["docx"])

    texts = " ".join(p.text for p in document.paragraphs)
    assert "The sitemap lists 0 addresses" not in texts
    assert "The sitemap could not be read" in texts

    rows = [[cell.text for cell in row.cells]
            for table in document.tables for row in table.rows]
    findings_named = [row[0] for row in rows]
    assert "The sitemap could not be read" in findings_named, findings_named
    assert "The robots file could not be read" in findings_named
    # And nothing claims the site answered everything.
    assert not any("every address the crawl asked for answered" in cell
                   for row in rows for cell in row)


def test_a_site_that_answers_reports_its_sitemap_as_before(tmp_path):
    sitemap_xml = (f'<?xml version="1.0"?><urlset>'
                   f'<url><loc>{BASE}/a</loc></url></urlset>')
    run = crawl_site(tmp_path, {BASE + "/": html_page(["/a"]),
                                BASE + "/a": html_page()},
                     sitemaps={BASE + "/sitemap.xml": sitemap_xml})
    findings = build_findings(run.out_dir, compare_enabled=False)

    assert run.summary["robots"]["refused"] is False
    assert run.summary["sitemap"]["refused"] == []
    assert findings["crawlability"]["sitemap"]["urls_total"] >= 1
    assert run.issues_of("request_refused") == []
