"""The report layer: guards, charts, document, validation. All offline."""

import json
import os

import pytest
import requests
import requests_mock

from seo_audit import narrative as narr
from seo_audit.findings import build_findings
from seo_audit.narrative import (API_URL, MAX_CALLS, NarrativeError, Narrator,
                                 check_numbers, strip_dashes, templated)
from seo_audit.report import (ValidationError, build_charts,
                              expected_headings_for, generate, pie_title,
                              validate)
from test_findings import make_run

HOST = "https://example.com"


def completion(text, prompt_tokens=100, completion_tokens=50):
    """A response shaped like the chat completions endpoint returns."""
    return {"choices": [{"message": {"role": "assistant", "content": text}}],
            "usage": {"prompt_tokens": prompt_tokens,
                      "completion_tokens": completion_tokens}}


@pytest.fixture
def run_dir(tmp_path):
    """A finished run folder with findings.json already built."""
    rows = [
        {"issue_type": "title_missing", "severity": "high",
         "url": f"{HOST}/p0", "final_url": f"{HOST}/p0",
         "detail": "no title", "site_level": "False"},
        {"issue_type": "orphan_page", "severity": "medium",
         "url": f"{HOST}/p1", "final_url": f"{HOST}/p1",
         "detail": "no crawled page links here", "site_level": "False"},
    ]
    ps = [{"url": f"{HOST}/p0", "template": "/p0", "group_size": "2",
           "strategy": s, "performance_score": "40", "lab_lcp_ms": "5000",
           "lab_cls": "0.3", "field_data_level": "url", "error": ""}
          for s in ("mobile", "desktop")]
    run = make_run(tmp_path, page_issues=rows, pagespeed=ps)
    build_findings(run, compare_enabled=False)
    return run


# --- guard 1: dashes ---------------------------------------------------------

def test_dashes_used_as_punctuation_are_replaced():
    assert "—" not in strip_dashes("The site is slow — this costs traffic.")
    assert "–" not in strip_dashes("Pages 10–20 were slow.")
    assert strip_dashes("It is slow - very slow.") == "It is slow, very slow."


def test_hyphens_inside_words_and_urls_survive():
    text = strip_dashes("Visit https://wework.co.in/on-demand/day-pass for "
                        "the on-demand plan.")
    assert "on-demand/day-pass" in text
    assert "the on-demand plan" in text


def test_a_replacement_never_leaves_a_stray_comma():
    assert strip_dashes("Slow — . Fix it.").count(",.") == 0
    assert ", ," not in strip_dashes("A — , B")


# --- guard 2: numbers --------------------------------------------------------

def test_a_sentence_with_an_invented_number_is_dropped():
    section = {"pages": {"count": 40, "whole": 842, "whole_is": "pages"}}
    text = ("We found 40 broken pages. Roughly 9999 other pages are at risk. "
            "That matters.")
    kept, dropped = check_numbers(text, section)

    assert "40 broken pages" in kept
    assert "That matters." in kept
    assert len(dropped) == 1
    assert "9999" in dropped[0]


def test_numbers_that_are_in_the_data_are_kept():
    section = {"a": 842, "b": 43.8, "c": {"count": 169, "whole": 842,
                                          "whole_is": "pages parsed"}}
    text = "We parsed 842 pages. The score was 43.8. That is 20% of them."
    kept, dropped = check_numbers(text, section)
    assert dropped == []
    assert kept.count(".") >= 3


def test_a_share_licenses_its_own_percentage():
    section = {"orphans": {"count": 169, "whole": 842,
                           "whole_is": "pages parsed"}}
    kept, dropped = check_numbers("About 20% of pages are orphans.", section)
    assert dropped == []
    assert "20%" in kept


def test_the_guard_logs_the_drop_and_the_finding_survives(tmp_path):
    """The number guard now runs on the model's half: why and what to do."""
    section = {"orphan_page": {"count": 40, "whole": 842,
                               "whole_is": "pages parsed"}}
    with requests_mock.Mocker() as mock:
        mock.post(API_URL, json=completion(
            '{"why": "Exactly 123456 pages are broken. It matters.", '
            '"todo": ["Fix it."]}'))
        narrator = Narrator(key="k", session=requests.Session())
        text = narrator.write("links", section)

    assert "123456" not in text["why"]
    assert "It matters." in text["why"]
    event = [e for e in narrator.usage.guard_events
             if "sentences_dropped" in e][0]
    assert event["section"] == "links"
    assert event["sentences_dropped"] == 1
    # The finding is stated by code, so it was never at risk.
    assert "40 of 842 pages have no links" in text["found"]


# --- calls, retries, refusals ------------------------------------------------

def test_a_rate_limit_is_retried_once(monkeypatch):
    monkeypatch.setattr(narr.time, "sleep", lambda s: None)
    with requests_mock.Mocker() as mock:
        mock.post(API_URL, [{"status_code": 429, "text": "slow down"},
                            {"status_code": 200,
                             "json": completion(
                                 '{"why": "All is well.", '
                                 '"todo": ["C."]}')}])
        narrator = Narrator(key="k", session=requests.Session())
        text = narrator.write("on_page", {"a": 1})

    assert "All is well." in text["why"]
    assert narrator.usage.calls == 2


def test_an_auth_refusal_stops_without_retry(monkeypatch):
    monkeypatch.setattr(narr.time, "sleep", lambda s: None)
    with requests_mock.Mocker() as mock:
        mock.post(API_URL, status_code=401, text="invalid api key")
        narrator = Narrator(key="k", session=requests.Session())
        with pytest.raises(NarrativeError, match="401"):
            narrator.write("on_page", {"a": 1})

    assert narrator.usage.calls == 1, "an auth failure must not be retried"


def test_an_insufficient_quota_body_stops_without_retry(monkeypatch):
    monkeypatch.setattr(narr.time, "sleep", lambda s: None)
    with requests_mock.Mocker() as mock:
        mock.post(API_URL, status_code=429,
                  text='{"error":{"code":"insufficient_quota"}}')
        narrator = Narrator(key="k", session=requests.Session())
        with pytest.raises(NarrativeError, match="insufficient_quota"):
            narrator.write("on_page", {"a": 1})
    assert narrator.usage.calls == 1


def test_the_call_ceiling_degrades_to_templates_rather_than_failing():
    """Out of budget must not lose the report: the rest is written from data."""
    narrator = Narrator(key="k", session=requests.Session())
    narrator.usage.calls = MAX_CALLS
    parts = narrator.write("on_page", {"a": 1})

    assert parts["found"], "the section still says something"
    assert any(e.get("reason") == "call ceiling reached"
               for e in narrator.usage.guard_events)
    assert narrator.usage.calls == MAX_CALLS, "no further calls were made"


def test_a_whole_report_never_exceeds_the_call_ceiling(run_dir):
    with requests_mock.Mocker() as mock:
        mock.post(API_URL, json=completion(
            '{"why": "It matters here.", "todo": ["Do the thing."]}'))
        result = generate(run_dir, use_ai=True, session=requests.Session())

    usage = result["usage_data"]
    assert usage["calls"] <= MAX_CALLS
    assert usage["input_tokens"] > 0 and usage["output_tokens"] > 0


# --- no-ai mode --------------------------------------------------------------

def test_no_ai_produces_a_complete_report_with_zero_calls(run_dir):
    with requests_mock.Mocker() as mock:
        mock.post(API_URL, status_code=500, text="must not be called")
        result = generate(run_dir, use_ai=False)
        assert mock.call_count == 0

    assert os.path.exists(result["docx"])
    assert result["usage_data"]["calls"] == 0
    assert result["usage_data"]["input_tokens"] == 0


def test_templated_prose_states_the_whole():
    text = templated("links", {"orphans": {"count": 169, "whole": 842,
                                           "whole_is": "pages parsed"}})
    assert "169 of 842 pages parsed" in text


# --- charts ------------------------------------------------------------------

def test_every_pie_title_states_its_whole(run_dir):
    with open(os.path.join(run_dir, "findings.json"), encoding="utf-8") as fh:
        findings = json.load(fh)
    charts = build_charts(findings)

    assert charts, "no charts were drawn"
    for key, (_buffer, title) in charts.items():
        if key.endswith("_bar"):
            continue
        assert " of " in title, f"{key}: {title}"
        assert any(ch.isdigit() for ch in title.split(" of ", 1)[1])


def test_pie_title_reads_as_words():
    assert pie_title("Page scores", 842, "pages parsed") == \
        "Page scores, of 842 pages parsed"


def test_no_png_files_are_left_behind(run_dir):
    generate(run_dir, use_ai=False)
    assert not [f for f in os.listdir(run_dir) if f.lower().endswith(".png")]


# --- the document ------------------------------------------------------------

def test_the_report_lands_next_to_the_run_with_its_usage_file(run_dir):
    result = generate(run_dir, use_ai=False)
    names = sorted(os.listdir(run_dir))

    assert os.path.basename(result["docx"]) in names
    assert "report_usage.json" in names
    assert os.path.basename(result["docx"]).startswith("SEO_Audit_")
    assert os.path.basename(result["docx"]).endswith(".docx")


def test_headings_appear_in_order(run_dir):
    from docx import Document
    result = generate(run_dir, use_ai=False)
    document = Document(result["docx"])
    headings = [p.text for p in document.paragraphs
                if p.style.name.startswith("Heading")
                or p.style.name == "Title"]

    with open(os.path.join(run_dir, "findings.json"), encoding="utf-8") as fh:
        expected = expected_headings_for(json.load(fh))
    position = 0
    for wanted in expected:
        assert wanted in headings[position:], f"{wanted} missing or misordered"
        position = headings.index(wanted, position) + 1


def test_the_document_carries_no_dash_used_as_punctuation(run_dir):
    from docx import Document
    result = generate(run_dir, use_ai=False)
    document = Document(result["docx"])

    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            texts.extend(cell.text for cell in row.cells)
    for text in texts:
        assert "—" not in text and "–" not in text
        assert " - " not in text


# --- comparison section ------------------------------------------------------

def test_the_comparison_section_is_absent_when_there_is_no_previous_run(
        run_dir):
    from docx import Document
    result = generate(run_dir, use_ai=False)
    headings = [p.text for p in Document(result["docx"]).paragraphs]
    assert "Comparison with the previous audit" not in headings


def test_a_coverage_change_is_stated_plainly(tmp_path):
    from docx import Document
    old = make_run(tmp_path, "20260101-000000", summary_extra={
        "sitemap_sweep": {"sitemap_urls_total": 106, "sitemap_urls_swept": 0,
                          "sweep_capped": False, "sweep_limit": 10000,
                          "sweep_throttled": False, "already_crawled": 106}})
    new = make_run(tmp_path, "20260102-000000", summary_extra={
        "sitemap_sweep": {"sitemap_urls_total": 709, "sitemap_urls_swept": 1,
                          "sweep_capped": False, "sweep_limit": 10000,
                          "sweep_throttled": False, "already_crawled": 708}})
    findings = build_findings(new, previous_dir=old)
    assert findings["comparison"]["coverage_changed"] is True

    result = generate(new, use_ai=False)
    document = Document(result["docx"])
    text = " ".join(p.text for p in document.paragraphs)
    assert "Comparison with the previous audit" in [
        p.text for p in document.paragraphs]
    assert "not comparable with each other" in text


# --- validation --------------------------------------------------------------

def test_validation_rejects_a_document_seeded_with_an_en_dash(run_dir,
                                                              tmp_path):
    from docx import Document
    result = generate(run_dir, use_ai=False)
    document = Document(result["docx"])
    document.paragraphs[2].text = "Traffic fell 10–20 percent this quarter."
    bad = str(tmp_path / "bad.docx")
    document.save(bad)

    with open(os.path.join(run_dir, "findings.json"), encoding="utf-8") as fh:
        findings = json.load(fh)
    problems = validate(bad, expected_headings_for(findings),
                        len(build_charts(findings)))
    assert any("dash" in p for p in problems)


def test_validation_rejects_a_missing_heading(run_dir, tmp_path):
    from docx import Document
    result = generate(run_dir, use_ai=False)
    document = Document(result["docx"])
    bad = str(tmp_path / "missing.docx")
    document.save(bad)

    problems = validate(bad, ["Executive summary", "A heading nobody wrote"],
                        len(build_charts(json.load(
                            open(os.path.join(run_dir, "findings.json"),
                                 encoding="utf-8")))))
    assert any("heading missing" in p for p in problems)


def test_validation_rejects_a_wrong_image_count(run_dir, tmp_path):
    result = generate(run_dir, use_ai=False)
    with open(os.path.join(run_dir, "findings.json"), encoding="utf-8") as fh:
        findings = json.load(fh)
    problems = validate(result["docx"], expected_headings_for(findings), 99)
    assert any("image count" in p for p in problems)


def test_a_valid_document_reports_no_problems(run_dir):
    result = generate(run_dir, use_ai=False)
    with open(os.path.join(run_dir, "findings.json"), encoding="utf-8") as fh:
        findings = json.load(fh)
    assert validate(result["docx"], expected_headings_for(findings),
                    len(build_charts(findings))) == []


def test_generate_raises_when_validation_fails(run_dir, monkeypatch):
    monkeypatch.setattr("seo_audit.report.validate",
                        lambda *a, **k: ["something is wrong"])
    with pytest.raises(ValidationError, match="something is wrong"):
        generate(run_dir, use_ai=False)


def test_a_run_folder_without_findings_is_refused(tmp_path):
    with pytest.raises(FileNotFoundError, match="findings.json"):
        generate(str(tmp_path), use_ai=False)


# --- the corrections A1 to A4 ------------------------------------------------

def test_a1_comparison_deltas_are_rounded(tmp_path):
    old = make_run(tmp_path, "20260101-000000", summary_extra={
        "pagespeed": {"skipped": False, "mean_mobile_score": 45.0,
                      "templates_sampled": 1, "pages_represented": 2,
                      "mean_desktop_score": 70.0, "errors": 0}})
    new = make_run(tmp_path, "20260102-000000", summary_extra={
        "pagespeed": {"skipped": False, "mean_mobile_score": 43.8,
                      "templates_sampled": 1, "pages_represented": 2,
                      "mean_desktop_score": 70.0, "errors": 0}})
    from seo_audit.findings import compare
    delta = compare(new, old)["deltas"]["mean_mobile_score"]["change"]
    assert delta == -1.2
    assert len(str(delta).split(".")[-1]) <= 1


def test_a2_one_entry_per_issue_type_with_siblings_merged(tmp_path):
    rows = []
    for section in ("alpha", "beta", "gamma"):
        for i in range(3):
            rows.append({"issue_type": "redirected_internal_link",
                         "severity": "low",
                         "url": f"{HOST}/{section}/p{i}",
                         "final_url": f"{HOST}/{section}/p{i}",
                         "detail": "linked from 100 page(s): a",
                         "site_level": "False"})
    findings = build_findings(make_run(tmp_path, page_issues=rows),
                              compare_enabled=False)
    fixes = findings["prioritised_fixes"]

    types = [f["issue_type"] for f in fixes]
    assert types.count("redirected_internal_link") == 1
    fix = fixes[0]
    assert fix["pattern_count"] == 3
    assert len(fix["sibling_patterns"]) == 2
    # A redirect target is a target, so the nine rows are nine targets and
    # the pages affected are the pages linking to them.
    assert fix["targets"]["count"] == 9
    assert fix["pages_affected"]["whole_is"] == "pages parsed"


def test_a3_impact_is_pages_times_severity_not_link_instances(tmp_path):
    rows = [
        # One low severity redirect group with an enormous link count.
        {"issue_type": "redirected_internal_link", "severity": "low",
         "url": f"{HOST}/a/p0", "final_url": f"{HOST}/a/p0",
         "detail": "linked from 9000 page(s): x", "site_level": "False"},
        # Three high severity pages with no link counts at all.
        *[{"issue_type": "title_missing", "severity": "high",
           "url": f"{HOST}/b/p{i}", "final_url": f"{HOST}/b/p{i}",
           "detail": "no title", "site_level": "False"} for i in range(3)],
    ]
    findings = build_findings(make_run(tmp_path, page_issues=rows),
                              compare_enabled=False)
    fixes = findings["prioritised_fixes"]

    assert fixes[0]["issue_type"] == "title_missing"
    assert fixes[0]["impact"] == 9          # 3 pages times weight 3
    redirect = [f for f in fixes
                if f["issue_type"] == "redirected_internal_link"][0]
    assert redirect["impact"] == 1          # 1 page times weight 1
    assert redirect["reach"] == 9000        # reach is recorded, not multiplied


def test_a4_a_coverage_change_marks_counts_not_comparable(tmp_path):
    old = make_run(tmp_path, "20260101-000000", summary_extra={
        "sitemap_sweep": {"sitemap_urls_total": 106, "sitemap_urls_swept": 0,
                          "sweep_capped": False, "sweep_limit": 10000,
                          "sweep_throttled": False, "already_crawled": 106},
        "links": {"pages_with_zero_inlinks": 56, "orphan_pages_in_sitemap": 56,
                  "broken_internal_targets": 0,
                  "redirected_internal_targets": 0,
                  "external_targets_found": 10, "external_checked": 10,
                  "external_broken": 1, "external_unchecked": 0,
                  "external_check_limit": 1000, "edges": 40}})
    new = make_run(tmp_path, "20260102-000000", summary_extra={
        "sitemap_sweep": {"sitemap_urls_total": 709, "sitemap_urls_swept": 1,
                          "sweep_capped": False, "sweep_limit": 10000,
                          "sweep_throttled": False, "already_crawled": 708},
        "links": {"pages_with_zero_inlinks": 169,
                  "orphan_pages_in_sitemap": 169,
                  "broken_internal_targets": 0,
                  "redirected_internal_targets": 0,
                  "external_targets_found": 10, "external_checked": 10,
                  "external_broken": 1, "external_unchecked": 0,
                  "external_check_limit": 1000, "edges": 40}})

    from seo_audit.findings import compare
    result = compare(new, old)
    assert result["coverage_changed"] is True

    orphans = [n for n in result["notable"] if n["metric"] == "orphan_pages"]
    assert orphans, "the orphan rise should still be listed"
    assert orphans[0]["note"] == "not comparable, coverage changed"
    sitemap = [n for n in result["notable"]
               if n["metric"] == "sitemap_urls_total"][0]
    # The sitemap is the site's own, so it is attributed to the site, not
    # written off as a coverage wobble.
    assert sitemap["attribution"] == "site"
    assert "the site's sitemap listed 106 URLs last time" in sitemap["note"]


# --- regressions caught by the first live report -----------------------------

def test_a_line_starting_with_a_hyphen_bullet_loses_the_marker():
    """Word folds the line break into a space, so "\n- x" resurfaces as " - x"."""
    text = "The URLs are:\n- https://a/faq/\n- https://b/x/"
    out = strip_dashes(text)

    assert " - " not in out
    assert not any(line.strip().startswith("-") for line in out.splitlines())
    # The items themselves survive, hyphens inside them included.
    assert "https://a/faq/" in out and "https://b/x/" in out


def test_asterisk_and_bullet_markers_go_the_same_way():
    assert " - " not in strip_dashes("Items:\n* one\n• two")
    assert "one" in strip_dashes("Items:\n* one\n• two")


def test_a_bulleted_list_survives_the_docx_round_trip(tmp_path):
    """The exact failure: guarded text, saved, reopened, still dash free."""
    from docx import Document
    document = Document()
    document.add_heading("Heading", level=1)
    document.add_paragraph(strip_dashes(
        "Five pages are marked noindex. The URLs are:\n"
        "- https://wework.co.in/faq/\n- https://wework.co.in/x/"))
    path = str(tmp_path / "round-trip.docx")
    document.save(path)

    reopened = Document(path)
    for paragraph in reopened.paragraphs:
        assert " - " not in paragraph.text


def test_a_template_name_in_braces_is_not_a_placeholder(run_dir, tmp_path):
    """"/workspaces/{slug}" is data this product produces, not an unfilled slot."""
    from docx import Document
    result = generate(run_dir, use_ai=False)
    document = Document(result["docx"])
    document.paragraphs[2].text = "The worst template is /workspaces/{slug}."
    path = str(tmp_path / "braces.docx")
    document.save(path)

    with open(os.path.join(run_dir, "findings.json"), encoding="utf-8") as fh:
        findings = json.load(fh)
    problems = validate(path, expected_headings_for(findings),
                        len(build_charts(findings)))
    assert not any("placeholder" in p for p in problems)


def test_a_real_unfilled_placeholder_is_still_caught(run_dir, tmp_path):
    from docx import Document
    result = generate(run_dir, use_ai=False)
    document = Document(result["docx"])
    document.paragraphs[2].text = "TODO write this section."
    path = str(tmp_path / "todo.docx")
    document.save(path)

    with open(os.path.join(run_dir, "findings.json"), encoding="utf-8") as fh:
        findings = json.load(fh)
    problems = validate(path, expected_headings_for(findings),
                        len(build_charts(findings)))
    assert any("placeholder" in p for p in problems)
