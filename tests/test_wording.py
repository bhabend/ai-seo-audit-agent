"""The last four wording defects, and the client read that follows them.

One fixture each: the sampling filter that ate a real instruction, American
spelling, template shapes in client sentences, and a findings paragraph that
had become a list.
"""

import json

import pytest
import requests
import requests_mock

from seo_audit import narrative as narr
from seo_audit.findings import build_findings, template_words
from seo_audit.narrative import (API_URL, Narrator, is_about_the_audit,
                                 section_facts, section_facts_parts,
                                 to_british)
from seo_audit.report import (build_charts, expected_headings_for, generate,
                              validate)
from test_facts import page_rows
from test_findings import make_run
from test_report import HOST, completion, run_dir  # noqa: F401  (fixture)


# --- item 1: the filter drops our measuring, not the client's verbs ---------

def test_the_filter_keeps_advice_that_merely_uses_the_word_audit():
    """"Audit and reduce third party scripts" is advice, not a note to us."""
    assert not is_about_the_audit(
        "Audit and reduce third party scripts on high traffic templates.")
    assert not is_about_the_audit(
        "Measure the effect of the change in your own analytics.")
    assert not is_about_the_audit("Crawl the site with your own tool.")

    assert is_about_the_audit(
        "Expand measurement to the 125 unmeasured templates.")
    assert is_about_the_audit("Re-run the audit after the fixes.")
    assert is_about_the_audit("Run the audit again next quarter.")
    assert is_about_the_audit("Crawl the site again after the fixes land.")


def test_a_useful_action_using_the_word_audit_reaches_the_document():
    body = ('{"why": "Speed decides whether people stay.", '
            '"todo": ["Audit and reduce third party scripts on high traffic '
            'templates."]}')
    with requests_mock.Mocker() as mock:
        mock.post(API_URL, json=completion(body))
        narrator = Narrator(key="k", session=requests.Session())
        parts = narrator.write("performance", {"a": 1})

    assert parts["todo"] == ["Audit and reduce third party scripts on high "
                             "traffic templates."]
    assert not [e for e in narrator.usage.guard_events
                if e.get("items_dropped")]


# --- item 2: British spelling, and words that say nothing -------------------

def test_american_spelling_is_corrected_rather_than_argued_with():
    assert to_british("optimized") == "optimised"
    assert to_british("Prioritize and organize") == "Prioritise and organise"
    assert to_british("summarizing the analyzed pages") == \
        "summarising the analysed pages"
    # Words that only look like the pattern are left alone.
    assert to_british("the size of the prize") == "the size of the prize"


def test_a_spelling_never_costs_a_second_call():
    body = ('{"why": "The pages are not optimized for phones.", '
            '"todo": ["Optimize the images on the workspace pages."]}')
    with requests_mock.Mocker() as mock:
        mock.post(API_URL, json=completion(body))
        narrator = Narrator(key="k", session=requests.Session())
        parts = narrator.write("performance", {"a": 1})

    assert narrator.usage.calls == 1
    assert "optimised" in parts["why"]
    assert parts["todo"][0].startswith("Optimise")


def test_an_empty_word_sends_the_section_back_once(monkeypatch):
    monkeypatch.setattr(narr.time, "sleep", lambda s: None)
    bad = completion('{"why": "Leverage the robust template system.", '
                     '"todo": ["Ensure the images are smaller."]}')
    good = completion('{"why": "Slow pages lose visitors.", '
                      '"todo": ["Make the images smaller."]}')
    with requests_mock.Mocker() as mock:
        mock.post(API_URL, [{"json": bad}, {"json": good}])
        narrator = Narrator(key="k", session=requests.Session())
        parts = narrator.write("performance", {"a": 1})

    assert narrator.usage.calls == 2
    assert parts["why"] == "Slow pages lose visitors."
    event = [e for e in narrator.usage.guard_events
             if "words the report does not use" in e["reason"]][0]
    assert "leverage" in event["reason"]


def test_the_style_asks_for_british_spelling_and_plain_words():
    assert "British spelling" in narr.STYLE
    for word in ("leverage", "robust", "ensure"):
        assert word in narr.STYLE
    assert "secure connection setting" in narr.STYLE


# --- item 3: template names in words ----------------------------------------

def test_a_template_shape_is_said_in_words():
    assert template_words("/workspaces/{slug} (depth 4)") == \
        "workspaces pages, 4 levels deep"
    assert template_words("/coworking-space/bangalore (depth 4)") == \
        "coworking space, bangalore pages, 4 levels deep"
    assert template_words("/ (homepage)") == "home page"
    assert template_words("/on-demand/{seg}?area") == \
        "on demand pages with a filter"


def worst_template_run(tmp_path):
    ps = [{"url": f"{HOST}/workspaces/a-long-slug-here", "strategy": s,
           "template": "/workspaces/{slug} (depth 4)", "group_size": "34",
           "performance_score": "15", "lab_lcp_ms": "9000", "lab_cls": "0.3",
           "field_data_level": "url", "error": ""}
          for s in ("mobile", "desktop")]
    return make_run(tmp_path, pagespeed=ps, summary_extra={
        "pagespeed": {"skipped": False, "templates_sampled": 1,
                      "pages_represented": 2, "mean_mobile_score": 15.0,
                      "mean_desktop_score": 30.0, "errors": 0,
                      "worst_template_mobile": {
                          "template": "/workspaces/{slug} (depth 4)",
                          "url": f"{HOST}/workspaces/a-long-slug-here",
                          "group_size": 34, "performance_score": 15}}})


def test_the_performance_section_names_templates_in_words(tmp_path):
    findings = build_findings(worst_template_run(tmp_path),
                              compare_enabled=False)
    performance = findings["performance"]

    # The shape stays in the data, for the tool and for the developers.
    assert performance["templates_measured"][0]["template"] == \
        "/workspaces/{slug} (depth 4)"
    assert performance["templates_measured"][0]["template_words"] == \
        "workspaces pages, 4 levels deep"

    facts = section_facts("performance", performance, 4)
    assert "workspaces pages, 4 levels deep" in facts
    assert "{" not in facts and "(depth" not in facts


def test_the_model_is_never_shown_a_shape(tmp_path):
    from seo_audit.narrative import curate

    findings = build_findings(worst_template_run(tmp_path),
                              compare_enabled=False)
    blob = json.dumps(curate(findings["performance"]))
    assert "{slug}" not in blob and "(depth" not in blob
    assert "workspaces pages, 4 levels deep" in blob


def test_a_template_shape_in_client_text_fails_validation(run_dir, tmp_path):
    from docx import Document

    result = generate(run_dir, use_ai=False)
    document = Document(result["docx"])
    document.paragraphs[2].text = "The slowest is /workspaces/{slug} (depth 4)."
    path = str(tmp_path / "shape.docx")
    document.save(path)

    with open(f"{run_dir}/findings.json", encoding="utf-8") as handle:
        findings = json.load(handle)
    problems = validate(path, expected_headings_for(findings),
                        len(build_charts(findings)))
    assert any("template shape" in p for p in problems)


def test_the_appendix_keeps_the_shape_for_the_developers(tmp_path):
    from docx import Document

    run = worst_template_run(tmp_path)
    build_findings(run, compare_enabled=False)
    result = generate(run, use_ai=False)

    document = Document(result["docx"])
    tables = [t for t in document.tables
              if [c.text for c in t.rows[0].cells][:2] == ["Template",
                                                           "Shape"]]
    assert tables, "no templates measured table"
    first = tables[0].rows[1].cells
    assert first[1].text == "/workspaces/{slug} (depth 4)"
    assert first[0].text == "workspaces pages, 4 levels deep"
    assert validate(result["docx"], expected_headings_for(
        json.load(open(f"{run}/findings.json", encoding="utf-8"))),
        len(build_charts(json.load(
            open(f"{run}/findings.json", encoding="utf-8"))))) == []


# --- item 4: six sentences, then a table ------------------------------------

def ten_finding_section():
    def s(count):
        return {"count": count, "whole": 842, "whole_is": "pages parsed"}
    return {
        "title": {"missing": s(3), "too_long": s(233), "too_short": s(14),
                  "duplicate_groups": 44},
        "meta_description": {"missing": s(9), "too_long": s(251),
                             "duplicate_groups": 50},
        "h1": {"missing": s(19), "multiple": s(72)},
        "images_missing_alt": s(280),
        "viewport_missing": s(0),
        "html_lang_missing": s(0),
    }


def sentences_in(text):
    import re
    return [s for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]


def test_ten_findings_become_six_sentences_and_a_table():
    paragraph, rows, everything = section_facts_parts(
        "on_page", ten_finding_section(), 842)

    assert len(sentences_in(paragraph)) == 6
    assert len(rows) == 4
    for _finding, count in rows:
        assert " of 842 pages parsed" in count or "groups of pages" in count
    # Nothing is lost and nothing is said twice.
    assert len(sentences_in(paragraph)) + len(rows) == 10
    for finding, _count in rows:
        assert finding.lower() not in paragraph.lower()
    assert len(sentences_in(everything)) == 10


def test_five_findings_need_no_table():
    section = {"title": {"missing": {"count": 3, "whole": 842,
                                     "whole_is": "pages parsed"}},
               "h1": {"missing": {"count": 19, "whole": 842,
                                  "whole_is": "pages parsed"}}}
    paragraph, rows, _everything = section_facts_parts("on_page", section, 842)
    assert rows == []
    assert len(sentences_in(paragraph)) == 2


def test_a_long_crawlability_section_also_holds_at_six():
    section = {
        "sitemap": {"urls_total": 709, "non_200": 3,
                    "in_sitemap_not_crawled": {"count": 1, "whole": 709,
                                               "whole_is": "sitemap URLs"},
                    "crawled_not_in_sitemap": {"count": 267, "whole": 948,
                                               "whole_is": "pages found"}},
        "redirects": {"total": {"count": 58, "whole": 948,
                                "whole_is": "pages found"},
                      "trailing_slash": {"count": 38, "whole": 58,
                                         "whole_is": "redirects"},
                      "chains": 4, "loops": 2},
        "robots": {"found": True, "blocked_linked": 0,
                   "blocked_in_sitemap": 1},
        "deep_pages": {"count": 17, "whole": 842, "whole_is": "pages parsed"},
        "render_suspects": {"count": 2, "whole": 842,
                            "whole_is": "pages parsed"},
        "slow_responses": 0, "fetch_errors": 0, "documents_linked": 20,
    }
    paragraph, rows, everything = section_facts_parts("crawlability", section)

    assert len(sentences_in(paragraph)) == 6
    assert rows, "the rest of the obstacles must still be reported"
    for finding, count in rows:
        assert finding and count
    assert "20 links" in [count for _f, count in rows]


def test_the_saved_report_carries_the_also_found_table(tmp_path):
    from docx import Document

    rows = []
    for issue_type, severity, count in (
            ("title_missing", "high", 3), ("title_too_long", "low", 20),
            ("title_too_short", "low", 4), ("h1_missing", "medium", 6),
            ("h1_multiple", "low", 7), ("images_missing_alt", "low", 8),
            ("viewport_missing", "medium", 5),
            ("html_lang_missing", "low", 9)):
        rows += page_rows(issue_type, severity, count,
                          shape="{host}/" + issue_type + "/p{i}")
    pages = [{"url": f"{HOST}/p{i}", "final_url": f"{HOST}/p{i}",
              "status_code": "200", "word_count": "300",
              "schema_block_count": "1", "has_microdata": "False",
              "canonical": f"{HOST}/p{i}", "score": "90"} for i in range(60)]
    counts = {}
    for row in rows:
        counts[row["issue_type"]] = counts.get(row["issue_type"], 0) + 1
    run = make_run(tmp_path, pages=pages, page_issues=rows,
                   summary_extra={"page_issue_counts": counts})
    build_findings(run, compare_enabled=False)
    result = generate(run, use_ai=False)

    document = Document(result["docx"])
    tables = [t for t in document.tables
              if [c.text for c in t.rows[0].cells] == ["Finding", "Count"]]
    assert tables, "the overflow findings never reached the document"
    assert any(p.text.strip() == "Also found."
               for p in document.paragraphs)
    for table in tables:
        for row in table.rows[1:]:
            finding, count = row.cells[0].text, row.cells[1].text
            assert finding and count
            assert any(ch.isdigit() for ch in count)


# --- the client read: a shape reached the appendix groups table -------------

def broken_links_run(tmp_path):
    rows = []
    for section, count in (("on-demand", 4), ("flexible-workspace", 2)):
        rows += [{"issue_type": "broken_internal_link", "severity": "high",
                  "url": f"{HOST}/{section}/city-{i}/?area=x",
                  "final_url": f"{HOST}/{section}/city-{i}/?area=x",
                  "detail": f"HTTP 404, linked from {10 + i} page(s): "
                            f"{HOST}/a, {HOST}/b",
                  "site_level": "False"} for i in range(count)]
    return make_run(tmp_path, page_issues=rows)


def test_the_broken_link_groups_are_named_in_words(tmp_path):
    """The client read found "/on-demand/{seg}?area" in the appendix."""
    from docx import Document

    run = broken_links_run(tmp_path)
    build_findings(run, compare_enabled=False)
    result = generate(run, use_ai=False)

    document = Document(result["docx"])
    tables = [t for t in document.tables
              if t.rows[0].cells[0].text == "Where they are"]
    assert tables, "no broken link groups table"
    for row in tables[0].rows[1:]:
        where = row.cells[0].text
        assert "{" not in where and "(depth" not in where, where
        assert where.endswith("pages") or "pages with a filter" in where


def test_only_the_shape_column_may_carry_a_shape(tmp_path):
    """The exemption is one column of one table, not every cell with a slash."""
    from docx import Document

    run = broken_links_run(tmp_path)
    findings = build_findings(run, compare_enabled=False)
    result = generate(run, use_ai=False)
    assert validate(result["docx"], expected_headings_for(findings),
                    len(build_charts(findings))) == []

    document = Document(result["docx"])
    table = [t for t in document.tables
             if t.rows[0].cells[0].text == "Where they are"][0]
    table.rows[1].cells[0].text = "/on-demand/{seg}?area"
    path = str(tmp_path / "leaked.docx")
    document.save(path)

    problems = validate(path, expected_headings_for(findings),
                        len(build_charts(findings)))
    assert any("template shape" in p for p in problems)
