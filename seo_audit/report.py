"""The deliverable: a validated Word document built from a finished run.

Two formats. The short one is what a client receives and is the default: a
score, a fixes table with an action on every row, one table per section, two
charts, and no prose beyond a sentence above each table. It is built entirely
from findings.json, so it calls no model and costs nothing. The long one,
kept for internal use behind --format long, is the narrated report the model
helps write.

Numbers come from code, prose comes from the model, and the document is
reopened and checked before this command will exit zero. A file that opens is
not a file that is right, so validation is part of producing it rather than
something a human does afterwards.

Every chart that shows a part of something is a pie whose title says what the
whole is, in words. That is the whole point of the share objects in
findings.json: a percentage with no stated denominator is how a capped sample
gets read as a fact about a site.

    python -m seo_audit.report --run output/example.com/20260101-120000
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
from datetime import date
from typing import Any, Dict, List, Optional, Tuple

from .findings import template_words
from .issues import label_of, unit_of
from .narrative import Narrator, strip_dashes

# Section order in the document. The key is the findings section it reads.
SECTIONS: List[Tuple[str, str]] = [
    ("crawlability", "Crawlability and sitemap"),
    ("indexability_technical", "Indexability and technical"),
    ("on_page", "On page"),
    ("content", "Content"),
    ("schema", "Structured data"),
    ("links", "Internal and external links"),
    ("performance", "Performance"),
]

CHART_WIDTH_INCHES = 5.8
MAX_DOC_BYTES = 2 * 1024 * 1024
TOP_FIXES = 10
APPENDIX_ROWS = 20

# Real unfilled placeholders. A bare "{" is not one: template names such as
# "/workspaces/{slug}" are data this product produces on purpose.
PLACEHOLDERS = ("TODO", "lorem ipsum", "{{", "}}", "{0}", "{}")

# Written by code, after the executive summary, so the narrative never has to
# stop and define a word mid sentence.
GLOSSARY = [
    ("Sitemap", "A file listing the pages a site wants search engines to "
     "know about."),
    ("Robots file", "A file at the root of a site telling search engines "
     "which addresses they may fetch."),
    ("Redirect", "An instruction sending a visitor from one address to "
     "another."),
    ("Preferred address", "The address a page names as the one it wants "
     "search engines to list, also called a canonical."),
    ("Hidden from search", "A page carrying an instruction that asks search "
     "engines not to list it, also called noindex."),
    ("Orphan page", "A page no other page on the site links to."),
    ("Internal link", "A link from one page of the site to another."),
    ("Structured data", "Extra machine readable notes in a page describing "
     "what it is about, used for richer search results."),
    ("Schema", "The shared vocabulary structured data is written in."),
    ("Language version", "An alternate page for another language or region, "
     "declared with hreflang."),
    ("Secure connection enforced", "A setting telling browsers to reach the "
     "site only over a secure connection, also called HSTS."),
    ("Core Web Vitals", "Google's three measures of how a page feels to "
     "use: loading, stability and responsiveness."),
    ("Largest contentful paint", "How long the main thing on the page takes "
     "to appear."),
    ("Cumulative layout shift", "How much the page moves about while it "
     "loads."),
    ("Interaction to next paint", "How quickly the page responds when "
     "someone taps or clicks."),
    ("Template", "A page layout shared by many pages, such as every blog "
     "post or every location page."),
    ("Page title", "The line a search engine shows as the clickable heading "
     "for a page, taken from the page's title."),
    ("Search description", "The short summary a page offers for search "
     "engines to show under its title, also called a meta description."),
]


# --- charts -----------------------------------------------------------------

def _pie(slices: List[Tuple[str, float]], title: str):
    """One pie, drawn to memory. Returns None when there is nothing to draw."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    slices = [(label, float(value)) for label, value in slices
              if value and float(value) > 0]
    if not slices:
        return None

    figure, axes = plt.subplots(figsize=(5.6, 3.6))
    values = [v for _l, v in slices]
    labels = [l for l, _v in slices]
    axes.pie(values, labels=labels, autopct="%1.0f%%", startangle=90,
             counterclock=False,
             textprops={"fontsize": 9})
    axes.axis("equal")
    axes.set_title(title, fontsize=10)
    buffer = io.BytesIO()
    figure.savefig(buffer, format="png", dpi=150, bbox_inches="tight")
    plt.close(figure)
    buffer.seek(0)
    return buffer


def _barh(rows: List[Tuple[str, float]], title: str, xlabel: str):
    """The one non share chart: template scores, which are not parts of a whole."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    rows = [(str(l), float(v)) for l, v in rows if v is not None][:12]
    if not rows:
        return None
    rows.reverse()

    figure, axes = plt.subplots(figsize=(6.0, max(2.4, 0.32 * len(rows) + 1)))
    axes.barh([l for l, _v in rows], [v for _l, v in rows])
    axes.set_xlim(0, 100)
    axes.set_xlabel(xlabel, fontsize=9)
    axes.set_title(title, fontsize=10)
    axes.tick_params(labelsize=8)
    buffer = io.BytesIO()
    figure.savefig(buffer, format="png", dpi=150, bbox_inches="tight")
    plt.close(figure)
    buffer.seek(0)
    return buffer


def pie_title(label: str, whole: int, whole_is: str) -> str:
    """Every pie says its whole, in words, so no share can float free."""
    return f"{label}, of {whole} {whole_is}"


def _sitemap_title(whole: int, crawled: int, not_crawled: int) -> str:
    """Spell the whole out: two different things are being added together."""
    word = "URL" if not_crawled == 1 else "URLs"
    return (f"Sitemap coverage, of {whole} URLs seen ({crawled} crawled plus "
            f"{not_crawled} sitemap {word} not crawled)")


def build_charts(findings: Dict) -> Dict[str, Any]:
    """Every chart the document can carry, keyed by the section it belongs to."""
    charts: Dict[str, Any] = {}

    headline = findings.get("headline") or {}
    distribution = headline.get("distribution") or {}
    scored = headline.get("pages_scored", 0)
    if distribution and scored:
        charts["headline"] = (
            _pie([(band, obj["count"]) for band, obj in distribution.items()],
                 pie_title("Page scores", scored, "pages scored")),
            pie_title("Page scores", scored, "pages scored"))

    schema = (findings.get("schema") or {}).get("coverage") or {}
    if schema:
        whole = schema.get("json_ld", {}).get("whole", 0)
        charts["schema"] = (
            _pie([("JSON-LD", schema.get("json_ld", {}).get("count", 0)),
                  ("Microdata only",
                   schema.get("microdata_only", {}).get("count", 0)),
                  ("No structured data",
                   schema.get("none", {}).get("count", 0))],
                 pie_title("Structured data coverage", whole, "pages parsed")),
            pie_title("Structured data coverage", whole, "pages parsed"))

    links = findings.get("links") or {}
    orphans = links.get("orphan_pages") or {}
    if orphans.get("whole"):
        whole = orphans["whole"]
        charts["links"] = (
            _pie([("Linked from other pages", whole - orphans["count"]),
                  ("No internal links in", orphans["count"])],
                 pie_title("Pages reachable by internal links", whole,
                           "pages parsed")),
            pie_title("Pages reachable by internal links", whole,
                      "pages parsed"))

    crawl = findings.get("crawlability") or {}
    sitemap = crawl.get("sitemap") or {}
    both = sitemap.get("crawled_not_in_sitemap") or {}
    not_crawled = sitemap.get("in_sitemap_not_crawled") or {}
    found = (findings.get("meta") or {}).get("coverage", {}).get("pages_found", 0)
    in_both = max(0, found - both.get("count", 0))
    whole = in_both + both.get("count", 0) + not_crawled.get("count", 0)
    if whole:
        charts["crawlability"] = (
            _pie([("In the sitemap and crawled", in_both),
                  ("Crawled, not in the sitemap", both.get("count", 0)),
                  ("In the sitemap, not crawled",
                   not_crawled.get("count", 0))],
                 _sitemap_title(whole, found, not_crawled.get("count", 0))),
            _sitemap_title(whole, found, not_crawled.get("count", 0)))

    redirects = crawl.get("redirects") or {}
    total = (redirects.get("total") or {}).get("count", 0)
    slash = (redirects.get("trailing_slash") or {}).get("count", 0)
    if total:
        charts["crawlability_redirects"] = (
            _pie([("Missing trailing slash", slash),
                  ("Other redirects", max(0, total - slash))],
                 pie_title("Redirects by kind", total, "redirects")),
            pie_title("Redirects by kind", total, "redirects"))

    performance = findings.get("performance") or {}
    if not performance.get("skipped"):
        represented = performance.get("pages_represented") or {}
        whole = represented.get("whole", 0)
        if whole:
            charts["performance"] = (
                _pie([("Covered by a measured template",
                       represented.get("count", 0)),
                      ("Not measured",
                       max(0, whole - represented.get("count", 0)))],
                     pie_title("Performance coverage", whole,
                               "pages parsed")),
                pie_title("Performance coverage", whole, "pages parsed"))
        measured = performance.get("templates_measured") or []
        charts["performance_bar"] = (
            _barh([(m.get("template_words") or m["template"],
                    m["mobile_score"]) for m in measured
                   if m.get("mobile_score") is not None],
                  "Mobile performance score by template",
                  "Score out of 100"),
            "Mobile performance score by template")
    return {k: v for k, v in charts.items() if v[0] is not None}


# --- document ---------------------------------------------------------------

def _add_table(document, columns: List[str], rows: List[List[str]]):
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Light Grid Accent 1"
    for index, name in enumerate(columns):
        table.rows[0].cells[index].text = strip_dashes(str(name))
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = strip_dashes(str(value))
    return table


def _add_section_body(document, parts: Dict) -> None:
    """What we found, why it matters, what to do, as three visible parts.

    Session 8 rendered each section as one wall of text and every "what to do"
    ran into the paragraph before it.
    """
    found = (parts.get("found") or "").strip()
    why = (parts.get("why") or "").strip()
    todo = [t for t in (parts.get("todo") or []) if str(t).strip()]

    if found:
        paragraph = document.add_paragraph()
        paragraph.add_run("What we found. ").bold = True
        paragraph.add_run(strip_dashes(found))
    also = [row for row in (parts.get("also_found") or []) if row]
    if also:
        # The findings that did not fit the paragraph. They are not lesser
        # findings, they are the ones a reader would rather scan than read.
        paragraph = document.add_paragraph()
        paragraph.add_run("Also found.").bold = True
        _add_table(document, ["Finding", "Count"],
                   [[str(row[0]), str(row[1])] for row in also])
    if why:
        paragraph = document.add_paragraph()
        paragraph.add_run("Why it matters. ").bold = True
        paragraph.add_run(strip_dashes(why))
    if todo:
        paragraph = document.add_paragraph()
        paragraph.add_run("What to do.").bold = True
        for item in todo:
            document.add_paragraph(strip_dashes(str(item)),
                                   style="List Number")


def _add_glossary(document) -> None:
    document.add_heading("Terms used in this report", level=1)
    document.add_paragraph(strip_dashes(
        "A few words appear throughout. They are set out here once so the "
        "sections that follow can get on with the findings."))
    for term, meaning in GLOSSARY:
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{term}. ").bold = True
        paragraph.add_run(strip_dashes(meaning))


def _add_chart(document, chart):
    from docx.shared import Inches
    buffer, _title = chart
    document.add_picture(buffer, width=Inches(CHART_WIDTH_INCHES))


def _add_comparison(document, comparison: Dict, narrator: Narrator) -> None:
    """When coverage moved, this section is written by code, not the model.

    A sitemap that went from 106 URLs to 709 is the site changing, not the
    crawler wobbling, and counts either side of that are not comparable. A
    model asked to narrate it will reach for a story about decline, so it is
    not asked.
    """
    sitemap = (comparison.get("deltas") or {}).get("sitemap_urls_total") or {}
    coverage_changed = comparison.get("coverage_changed")

    if coverage_changed:
        if comparison.get("sitemap_changed"):
            document.add_paragraph(strip_dashes(
                f"The site's own sitemap changed between the two audits. It "
                f"listed {sitemap.get('previous')} addresses last time and "
                f"{sitemap.get('now')} this time. That is a change on the "
                f"site, not in how it was audited."))
        document.add_paragraph(strip_dashes(
            "Because the two audits saw different amounts of the site, the "
            "counts below are not comparable with each other. A number that "
            "looks higher or lower mostly reflects how many pages were "
            "reached, so this report draws no conclusion from it."))
        document.add_paragraph(strip_dashes(
            "Worth checking on your side: a sitemap that reports a different "
            "size from one fetch to the next usually points to a caching or "
            "generation problem, and it affects what search engines see too."))
    else:
        _add_section_body(document, narrator.write("comparison", comparison))

    notable = comparison.get("notable") or []
    if notable:
        _add_table(document,
                   ["Measure", "Previous", "Now", "Change", "Note"],
                   [[_metric_label(n["metric"]), n["previous"], n["now"],
                     n["change"], n.get("note", "")]
                    for n in notable[:APPENDIX_ROWS]])


def _metric_label(metric: str) -> str:
    """Plain names for the handful of metrics the comparison tracks."""
    names = {
        "pages_found": "addresses crawled",
        "pages_parsed": "pages read",
        "sitemap_urls_total": "addresses in the sitemap",
        "site_score": "overall score",
        "mean_mobile_score": "average mobile speed score",
        "orphan_pages": "pages with no links in",
        "broken_internal_targets": "broken internal addresses",
        "external_broken": "broken links to other sites",
    }
    return names.get(metric, metric.replace("_", " "))


def _coverage_paragraph(findings: Dict) -> str:
    meta = findings.get("meta") or {}
    coverage = meta.get("coverage") or {}
    caps = meta.get("caps_applied") or {}
    performance = findings.get("performance") or {}
    templates = performance.get("templates_sampled", 0)

    sentence = (
        f"This audit crawled {coverage.get('pages_found', 0)} URLs on "
        f"{meta.get('host', 'the site')} and read "
        f"{coverage.get('pages_parsed', 0)} of them as pages. The sitemap "
        f"listed {coverage.get('sitemap_urls_total', 0)} URLs. "
        f"The crawl stopped at {caps.get('max_pages')} pages and "
        f"{caps.get('max_depth')} clicks from the home page, checked at most "
        f"{caps.get('external_check_limit')} links to other websites, and "
        f"measured speed on {templates} page templates rather than every "
        f"page, because pages built from one template perform alike. "
        f"Every figure in this report says what it is a share of, so a "
        f"sampled number is never read as a total.")
    return strip_dashes(sentence)


SCOPE_WORDS = {"config": "one setting", "template": "one template",
               "page": "page by page"}


def _fixes_table_rows(fixes: List[Dict]) -> List[List[str]]:
    rows = []
    for fix in fixes[:TOP_FIXES]:
        affected = fix.get("pages_affected") or {}
        pages = (f"{affected.get('count', 0)} of {affected.get('whole', 0)} "
                 f"{affected.get('whole_is', 'pages')}")
        targets = fix.get("targets")
        if targets:
            # "118 of 842 pages" was 118 external links, and the pages are
            # the pages linking to them. The broken things come first and
            # the pages follow, so neither number can be read as the other.
            scale = (f"{targets['count']} of {targets['whole']} "
                     f"{targets['whole_is']}, linked from {pages}")
        else:
            scale = pages
        rows.append([
            fix.get("title", ""),
            scale,
            fix.get("severity", ""),
            SCOPE_WORDS.get(fix.get("fix_scope", ""), fix.get("fix_scope", "")),
        ])
    return rows


def build_document(findings: Dict, narrator: Narrator, charts: Dict,
                   host: str) -> Any:
    from docx import Document

    document = Document()
    today = date.today().isoformat()

    document.add_heading(f"SEO Audit: {host}", level=0)
    document.add_paragraph(strip_dashes(
        f"Prepared {today}. Mode: "
        f"{(findings.get('meta') or {}).get('mode', 'baseline')}."))
    document.add_paragraph(_coverage_paragraph(findings))

    document.add_heading("Executive summary", level=1)
    _add_section_body(document, narrator.summary(
        findings.get("headline") or {},
        findings.get("prioritised_fixes") or []))
    if "headline" in charts:
        _add_chart(document, charts["headline"])

    _add_glossary(document)

    document.add_heading("Priority fixes", level=1)
    fixes = findings.get("prioritised_fixes") or []
    if fixes:
        _add_table(document, ["Fix", "Pages affected", "Severity", "Scope"],
                   _fixes_table_rows(fixes))
    _add_section_body(document, narrator.write("prioritised_fixes",
                                               fixes[:TOP_FIXES]))

    for key, heading in SECTIONS:
        section = findings.get(key)
        if section is None:
            continue
        document.add_heading(heading, level=1)
        _add_section_body(document, narrator.write(key, section))
        if key in charts:
            _add_chart(document, charts[key])
        if key == "crawlability" and "crawlability_redirects" in charts:
            _add_chart(document, charts["crawlability_redirects"])
        if key == "performance" and "performance_bar" in charts:
            _add_chart(document, charts["performance_bar"])

    comparison = findings.get("comparison") or {}
    if comparison.get("previous_run"):
        document.add_heading("Comparison with the previous audit", level=1)
        _add_comparison(document, comparison, narrator)

    document.add_heading("Search performance", level=1)
    document.add_paragraph(strip_dashes(
        "Search Console data was not provided for this audit, so this report "
        "does not cover what people searched for, how often the site "
        "appeared, or how often it was clicked. Supply a Search Console "
        "export and a later audit can include it."))

    document.add_heading("Appendix", level=1)
    links = findings.get("links") or {}
    broken = (links.get("broken_internal") or {}).get("groups") or []
    if broken:
        document.add_heading("Broken internal link groups", level=2)
        _add_table(document,
                   ["Where they are", "Addresses", "Links pointing there"],
                   [[template_words(g["pattern"]), g["targets"],
                     g["link_instances"]]
                    for g in broken[:APPENDIX_ROWS]])

    lowest = (findings.get("headline") or {}).get("lowest_pages") or []
    if lowest:
        document.add_heading("Lowest scoring pages", level=2)
        _add_table(document, ["Page", "Score", "Main issue"],
                   [[p["final_url"], p["score"],
                     label_of(p["top_issue"]) if p.get("top_issue") else ""]
                    for p in lowest[:APPENDIX_ROWS]])

    measured = (findings.get("performance") or {}).get("templates_measured")
    if measured:
        document.add_heading("Templates measured for speed", level=2)
        document.add_paragraph(strip_dashes(
            "The shape column is how the audit tool names each template. It "
            "is here for whoever maintains the site."))
        _add_table(document,
                   ["Template", "Shape", "Pages", "Mobile", "Desktop"],
                   [[m.get("template_words") or m["template"], m["template"],
                     m["group_size"], m["mobile_score"], m["desktop_score"]]
                    for m in measured[:APPENDIX_ROWS]])
    return document


# --- the short report: tables, not prose ------------------------------------
#
# Management's rule for what a client receives: short, and not obviously
# written by a machine. So nothing here is written by a machine that writes
# sentences. Every line is built from findings.json, and every finding
# arrives as a row with its count, its whole, its severity and the one thing
# to do about it. No model is called, so a short report costs nothing.

SHORT_MAX_WORDS = 2500
SHORT_IMAGES = 2
SHORT_SECTION_COLUMNS = ["Finding", "Count", "Severity", "Action"]
FIXES_COLUMNS = ["Rank", "Fix", "Pages affected", "Severity", "Scope",
                 "Action"]
IN_PLACE = "In place"
SEVERITY_RANK = {"high": 0, "medium": 1, "low": 2, "unmeasured": 3, "info": 4}


# findings.json names its wholes for the tool. These are the same wholes in
# the words the rest of the report uses.
WHOLE_WORDS = {
    "pages found": "addresses crawled",
    "sitemap URLs": "sitemap addresses",
    "external targets found": "links to other websites found",
    "external targets checked": "links to other websites checked",
    "internal addresses crawled": "addresses crawled",
    "canonical check limit": "addresses checked",
}


def _plain_count(count: int, whole: int, whole_is: str) -> str:
    return f"{count} of {whole} {WHOLE_WORDS.get(whole_is, whole_is)}".strip()


def _count_cell(issue_type: str, count: int, whole: int, whole_is: str,
                parsed: int, crawled: int) -> str:
    """A count that says what it is out of, whatever it counts.

    A group of pages is not a page and a broken address is not a page, so
    each unit gets its own whole rather than borrowing the page total.
    """
    unit = unit_of(issue_type)
    if unit == "site":
        return "the whole site"
    if whole:
        return _plain_count(count, whole, whole_is)
    if unit == "group":
        return f"{count} groups of {parsed} pages parsed"
    if unit == "target":
        return _plain_count(count, crawled, "addresses crawled")
    return _plain_count(count, parsed, "pages parsed")


def _section_table(section_name: str, section: Any, parsed: int,
                   crawled: int) -> Tuple[List[List[str]], List[str]]:
    """(rows, what is already in place) for one findings section."""
    from .issues import PAGE_ISSUE_ACTION, PAGE_ISSUE_SEVERITY
    from .narrative import FACT_PATHS, fact_entries

    entries = fact_entries(section_name, section, parsed)
    found: List[Tuple[int, int, List[str]]] = []
    in_place: List[str] = []
    named = set()
    for issue_type, count, whole, whole_is in entries:
        named.add(issue_type)
        severity = PAGE_ISSUE_SEVERITY.get(issue_type, "low")
        cell = _count_cell(issue_type, count, whole, whole_is, parsed,
                           crawled)
        if severity == "info":
            in_place.append(f"{label_of(issue_type)}, {cell}")
            continue
        found.append((SEVERITY_RANK.get(severity, 2), -count,
                      [label_of(issue_type).capitalize(), cell, severity,
                       PAGE_ISSUE_ACTION.get(issue_type, "")]))

    if section_name == "content":
        handled = sum((section.get(key) or {}).get("canonical_handled", 0)
                      for key in ("duplicate_content",
                                  "near_duplicate_content"))
        if handled:
            in_place.append(f"{handled} groups of pages with the same text "
                            f"already name one of themselves as the "
                            f"preferred address")

    # What the section checked and did not find. A client reads this as the
    # part of their site that is already right, which is why it is here.
    clean = [label_of(issue_type)
             for issue_type, _path in FACT_PATHS.get(section_name, ())
             if issue_type not in named]
    if clean:
        in_place.append("nothing found for " + ", ".join(clean[:4]))

    found.sort(key=lambda item: (item[0], item[1]))
    return [row for _rank, _count, row in found], in_place


def _in_place_row(in_place: List[str]) -> List[str]:
    text = "; ".join(in_place) if in_place else "nothing to note"
    return [IN_PLACE, "", "", text[0].upper() + text[1:] + "."]


def short_section_parts(key: str, section: Any, coverage: Dict
                        ) -> Tuple[List[List[str]], List[str]]:
    """(finding rows, the In place row) for one section of the short report.

    The one place a section's table is decided. The document renders it and
    the console shows the same rows, so an operator and a client are never
    looking at two different tables built by two different pieces of code.
    """
    parsed = coverage.get("pages_parsed", 0)
    crawled = coverage.get("pages_found", 0)
    if key == "crawlability":
        rows, in_place = _crawlability_table(section, coverage)
    elif key == "performance":
        rows, in_place = _performance_table(section)
    else:
        rows, in_place = _section_table(key, section, parsed, crawled)
    return rows, _in_place_row(in_place)


def short_section_rows(key: str, section: Any, coverage: Dict
                       ) -> List[List[str]]:
    """The whole table, findings first and what is in place last."""
    rows, in_place_row = short_section_parts(key, section, coverage)
    return rows + [in_place_row]


# Crawlability counts obstacles rather than page faults, so its rows are
# written here: what was in the way, how many of what, and what to do.
def _crawlability_table(section: Dict, coverage: Dict
                        ) -> Tuple[List[List[str]], List[str]]:
    sitemap = section.get("sitemap") or {}
    redirects = section.get("redirects") or {}
    robots = section.get("robots") or {}
    crawled = coverage.get("pages_found", 0)
    listed = sitemap.get("urls_total", 0)
    redirected = (redirects.get("total") or {}).get("count", 0)

    def share(obj) -> Tuple[int, str]:
        obj = obj or {}
        return (obj.get("count", 0),
                _plain_count(obj.get("count", 0), obj.get("whole", 0),
                             obj.get("whole_is", "")))

    rows: List[Tuple[int, int, List[str]]] = []

    def add(count: int, cell: str, finding: str, severity: str, action: str):
        if count:
            rows.append((SEVERITY_RANK.get(severity, 2), -count,
                         [finding, cell, severity, action]))

    count, cell = share(sitemap.get("in_sitemap_not_crawled"))
    add(count, cell, "Sitemap addresses the crawl never reached", "medium",
        "Link these addresses from the site, or take them out of the "
        "sitemap.")
    count, cell = share(sitemap.get("crawled_not_in_sitemap"))
    add(count, cell, "Addresses found by following links but not in the "
        "sitemap", "medium", "Add these addresses to the sitemap.")
    add(sitemap.get("non_200", 0),
        _plain_count(sitemap.get("non_200", 0), listed, "sitemap addresses"),
        "Sitemap entries that do not load", "high",
        "Repair these pages or take them out of the sitemap.")

    count, cell = share(redirects.get("total"))
    add(count, cell, "Addresses that answer with a redirect", "low",
        "Point internal links at the address the redirect ends on.")
    count, cell = share(redirects.get("trailing_slash"))
    add(count, cell, "Redirects that only add a slash at the end", "low",
        "Write internal links with the slash at the end so they land "
        "directly.")
    add(redirects.get("chains", 0),
        _plain_count(redirects.get("chains", 0), redirected, "redirects"),
        "Redirects that pass through another redirect", "low",
        "Point these addresses straight at the final page.")
    add(redirects.get("loops", 0),
        _plain_count(redirects.get("loops", 0), redirected, "redirects"),
        "Redirect loops", "high",
        "Repair these addresses so they arrive at a page.")

    add(robots.get("blocked_linked", 0),
        _plain_count(robots.get("blocked_linked", 0), crawled,
                     "addresses crawled"),
        "Linked addresses the robots file blocks", "high",
        "Allow these addresses in the robots file, or stop linking to them.")
    add(robots.get("blocked_in_sitemap", 0),
        _plain_count(robots.get("blocked_in_sitemap", 0), listed,
                     "sitemap addresses"),
        "Sitemap addresses the robots file blocks", "high",
        "Allow these addresses in the robots file, or take them out of the "
        "sitemap.")

    count, cell = share(section.get("deep_pages"))
    add(count, cell, "Pages more than three clicks from the home page", "low",
        "Link these pages from a navigation or category page.")
    count, cell = share(section.get("render_suspects"))
    add(count, cell, "Pages with almost no text until scripts run", "medium",
        "Serve the main text of these pages in the page itself.")
    add(section.get("slow_responses", 0),
        _plain_count(section.get("slow_responses", 0), crawled,
                     "addresses crawled"),
        "Addresses slow to answer", "medium",
        "Ask the developers what makes these addresses slow.")
    add(section.get("fetch_errors", 0),
        _plain_count(section.get("fetch_errors", 0), crawled,
                     "addresses crawled"),
        "Addresses that returned nothing", "high",
        "Repair or remove these addresses.")
    add(section.get("documents_linked", 0),
        _plain_count(section.get("documents_linked", 0), crawled,
                     "addresses crawled"),
        "Links to documents rather than pages", "low",
        "Link to a page describing the document where one would serve "
        "better.")
    add(section.get("non_html_linked", 0),
        _plain_count(section.get("non_html_linked", 0), crawled,
                     "addresses crawled"),
        "Links to files that are neither pages nor documents", "low",
        "Point these links at a page, or remove them.")

    in_place = []
    if robots.get("found"):
        in_place.append("a robots file is in place")
    if listed:
        in_place.append(f"a sitemap is in place, listing {listed} addresses")
    if not section.get("fetch_errors"):
        in_place.append("every address the crawl asked for answered")

    rows.sort(key=lambda item: (item[0], item[1]))
    return [row for _rank, _count, row in rows], in_place


def _performance_table(section: Dict) -> Tuple[List[List[str]], List[str]]:
    """Speed is measured per template, so the rows count templates."""
    from .issues import PAGE_ISSUE_ACTION, PAGE_ISSUE_SEVERITY
    from .pagespeed import POOR_CLS, POOR_LCP_MS, POOR_PERFORMANCE_SCORE

    measured = section.get("templates_measured") or []
    total = len(measured)
    if not total:
        return [], ["speed was not measured for this audit"]

    def count_where(test) -> int:
        return sum(1 for m in measured if test(m))

    slow = count_where(lambda m: m.get("mobile_score") is not None
                       and m["mobile_score"] < POOR_PERFORMANCE_SCORE)
    late = count_where(lambda m: (m.get("lab_lcp_ms") or 0) > POOR_LCP_MS)
    shifty = count_where(lambda m: (m.get("lab_cls") or 0) > POOR_CLS)

    rows = []
    for count, issue_type, finding in (
            (slow, "performance_poor",
             "Templates that score poorly on a phone"),
            (late, "lcp_poor",
             "Templates whose main content takes a long time to appear"),
            (shifty, "cls_poor",
             "Templates that move about while they load")):
        if count:
            rows.append([finding,
                         f"{count} of {total} measured templates",
                         PAGE_ISSUE_SEVERITY.get(issue_type, "medium"),
                         PAGE_ISSUE_ACTION.get(issue_type, "")])

    healthy = total - slow
    in_place = []
    if healthy:
        in_place.append(f"{healthy} of {total} measured templates score 50 "
                        f"or better on a phone")
    return rows, in_place


def _short_lead(section_name: str, section: Any, rows: List[List[str]],
                coverage: Dict) -> str:
    """One sentence above each table, carrying the number that matters most."""
    if section_name == "performance":
        measured = section.get("templates_measured") or []
        represented = section.get("pages_represented") or {}
        if not measured:
            return "Speed was not measured for this audit."
        return (f"Speed was measured on {len(measured)} page templates "
                f"covering {represented.get('count', 0)} of "
                f"{represented.get('whole', 0)} pages read, averaging "
                f"{section.get('mean_mobile_score')} of 100 on a phone and "
                f"{section.get('mean_desktop_score')} of 100 on a desktop.")
    if section_name == "crawlability":
        sitemap = (section or {}).get("sitemap") or {}
        return (f"The crawl reached {coverage.get('pages_found', 0)} "
                f"addresses and the sitemap lists "
                f"{sitemap.get('urls_total', 0)}.")
    if not rows:
        return "Nothing to fix in this section."
    first = rows[0][0]
    return (f"{len(rows)} findings here, the most serious first: "
            f"{first[0].lower() + first[1:]}, {rows[0][1]}.")


def _short_fixes_rows(fixes: List[Dict]) -> List[List[str]]:
    from .issues import PAGE_ISSUE_ACTION

    rows = []
    for rank, fix in enumerate(fixes[:TOP_FIXES], start=1):
        affected = fix.get("pages_affected") or {}
        scale = _plain_count(affected.get("count", 0),
                             affected.get("whole", 0),
                             affected.get("whole_is", "pages"))
        targets = fix.get("targets")
        if targets:
            scale += (f" ({targets['count']} of {targets['whole']} "
                      f"{targets['whole_is']})")
        label = fix.get("label") or fix.get("title", "")
        rows.append([
            str(rank),
            label[0].upper() + label[1:] if label else "",
            scale,
            fix.get("severity", ""),
            SCOPE_WORDS.get(fix.get("fix_scope", ""), fix.get("fix_scope", "")),
            PAGE_ISSUE_ACTION.get(fix.get("issue_type", ""), ""),
        ])
    return rows


def build_short_charts(findings: Dict) -> Dict[str, Any]:
    """Two charts: how the pages score, and how the templates perform."""
    charts: Dict[str, Any] = {}
    headline = findings.get("headline") or {}
    distribution = headline.get("distribution") or {}
    scored = headline.get("pages_scored", 0)
    if distribution and scored:
        title = pie_title("Page scores", scored, "pages scored")
        charts["headline"] = (
            _pie([(band.replace("-", " to "), obj["count"])
                  for band, obj in distribution.items()], title), title)

    measured = ((findings.get("performance") or {})
                .get("templates_measured") or [])
    rows = [(m.get("template_words") or m["template"], m["mobile_score"])
            for m in measured if m.get("mobile_score") is not None]
    if rows:
        title = "Mobile speed score by template, out of 100"
        charts["performance_bar"] = (
            _barh(rows, title, "Score out of 100"), title)
    return {key: value for key, value in charts.items() if value[0] is not None}


def _short_cover(document, findings: Dict, host: str) -> None:
    meta = findings.get("meta") or {}
    coverage = meta.get("coverage") or {}
    caps = meta.get("caps_applied") or {}
    templates = (findings.get("performance") or {}).get("templates_sampled", 0)

    document.add_heading(f"SEO Audit: {host}", level=0)
    document.add_paragraph(strip_dashes(
        f"Prepared {date.today().isoformat()}. Mode: "
        f"{meta.get('mode', 'baseline')}."))
    document.add_paragraph(strip_dashes(
        f"The crawl reached {coverage.get('pages_found', 0)} addresses on "
        f"{host} and read {coverage.get('pages_parsed', 0)} of them as "
        f"pages. The sitemap lists {coverage.get('sitemap_urls_total', 0)} "
        f"addresses."))
    document.add_paragraph(strip_dashes(
        f"Limits: at most {caps.get('max_pages')} pages, "
        f"{caps.get('max_depth')} clicks from the home page, "
        f"{caps.get('external_check_limit')} links to other websites "
        f"checked, and speed measured on {templates} page templates."))
    document.add_paragraph(strip_dashes(
        "Every count below says what it is a share of."))


def _short_score(document, findings: Dict, charts: Dict) -> None:
    headline = findings.get("headline") or {}
    deduction = headline.get("site_level_deduction") or 0
    excluded = (headline.get("noindex_excluded") or {}).get("count", 0)

    document.add_heading("Score", level=1)
    sentence = (f"The site scores {headline.get('site_score')} out of 100 "
                f"across {headline.get('pages_scored', 0)} pages scored")
    if deduction:
        sentence += (f", after a deduction of {deduction} points for "
                     f"settings that apply to the whole site")
    if excluded:
        sentence += (f", with {excluded} pages left out because they are "
                     f"hidden from search")
    document.add_paragraph(strip_dashes(sentence + "."))
    if "headline" in charts:
        _add_chart(document, charts["headline"])


def build_short_document(findings: Dict, charts: Dict, host: str) -> Any:
    """The whole client deliverable, built from findings.json alone."""
    from docx import Document

    document = Document()
    coverage = (findings.get("meta") or {}).get("coverage") or {}
    parsed = coverage.get("pages_parsed", 0)
    crawled = coverage.get("pages_found", 0)

    _short_cover(document, findings, host)
    _short_score(document, findings, charts)

    document.add_heading("Priority fixes", level=1)
    fixes = findings.get("prioritised_fixes") or []
    if fixes:
        document.add_paragraph(strip_dashes(
            f"The {min(len(fixes), TOP_FIXES)} fixes worth doing first, "
            f"ordered by how much of the site they affect and how serious "
            f"they are."))
        _add_table(document, FIXES_COLUMNS, _short_fixes_rows(fixes))
    else:
        document.add_paragraph("No fixes were raised by this audit.")

    for key, heading in SECTIONS:
        section = findings.get(key)
        if section is None:
            continue
        rows, in_place_row = short_section_parts(key, section, coverage)

        document.add_heading(heading, level=1)
        document.add_paragraph(strip_dashes(
            _short_lead(key, section, rows, coverage)))
        _add_table(document, SHORT_SECTION_COLUMNS, rows + [in_place_row])

        if key == "performance":
            measured = (section or {}).get("templates_measured") or []
            if measured:
                _add_table(document,
                           ["Template", "Shape for developers", "Pages",
                            "Mobile", "Desktop"],
                           [[m.get("template_words") or m["template"],
                             m["template"], m["group_size"],
                             m["mobile_score"], m["desktop_score"]]
                            for m in measured[:APPENDIX_ROWS]])
            if "performance_bar" in charts:
                _add_chart(document, charts["performance_bar"])

    comparison = findings.get("comparison") or {}
    if comparison.get("previous_run"):
        document.add_heading("Comparison with the previous audit", level=1)
        _short_comparison(document, comparison)

    document.add_heading("Search performance", level=1)
    document.add_paragraph(strip_dashes(
        "Search Console data was not supplied, so this report does not cover "
        "what people searched for or how often the site was clicked."))

    document.add_heading("Appendix", level=1)
    links = findings.get("links") or {}
    broken = (links.get("broken_internal") or {}).get("groups") or []
    if broken:
        document.add_heading("Broken internal link groups", level=2)
        _add_table(document,
                   ["Where they are", "Addresses", "Links pointing there"],
                   [[template_words(g["pattern"]), g["targets"],
                     g["link_instances"]]
                    for g in broken[:APPENDIX_ROWS]])
    lowest = (findings.get("headline") or {}).get("lowest_pages") or []
    if lowest:
        document.add_heading("Lowest scoring pages", level=2)
        _add_table(document, ["Page", "Score", "Main issue"],
                   [[p["final_url"], p["score"],
                     label_of(p["top_issue"]) if p.get("top_issue") else ""]
                    for p in lowest[:10]])
    return document


def _short_comparison(document, comparison: Dict) -> None:
    """The same attribution the long report makes, in two sentences."""
    sitemap = (comparison.get("deltas") or {}).get("sitemap_urls_total") or {}
    if comparison.get("coverage_changed"):
        if comparison.get("sitemap_changed"):
            document.add_paragraph(strip_dashes(
                f"The site's own sitemap changed between the two audits, "
                f"from {sitemap.get('previous')} addresses to "
                f"{sitemap.get('now')}. That is a change on the site, not in "
                f"how it was audited."))
        document.add_paragraph(strip_dashes(
            "The two audits saw different amounts of the site, so the counts "
            "below are not comparable."))
    else:
        document.add_paragraph(strip_dashes(
            "Both audits saw the same amount of the site, so the counts "
            "below are comparable."))
    notable = comparison.get("notable") or []
    if notable:
        _add_table(document, ["Measure", "Previous", "Now", "Change", "Note"],
                   [[_metric_label(n["metric"]), n["previous"], n["now"],
                     n["change"], n.get("note", "")]
                    for n in notable[:APPENDIX_ROWS]])


def expected_headings_for_short(findings: Dict) -> List[str]:
    headings = ["Score", "Priority fixes"]
    headings += [title for key, title in SECTIONS
                 if findings.get(key) is not None]
    if (findings.get("comparison") or {}).get("previous_run"):
        headings.append("Comparison with the previous audit")
    headings += ["Search performance", "Appendix"]
    return headings


# --- validation -------------------------------------------------------------

class ValidationError(RuntimeError):
    """The document was written but is not fit to send."""


_BAD_DASH = re.compile(r"[‐-―−]|(?<=\S) - (?=\S)")
# A finished sentence. An action item that does not match this stopped early.
_ENDS_PROPERLY = re.compile(r"[.!?][\"')\]]?$")


def _sections_missing_parts(document) -> List[str]:
    """Narrative headings whose body lacks found, why or what to do."""
    narrative_headings = {title for _key, title in SECTIONS}
    narrative_headings.add("Executive summary")
    narrative_headings.add("Priority fixes")

    missing: List[str] = []
    current: Optional[str] = None
    seen: set = set()
    for paragraph in document.paragraphs:
        style = paragraph.style.name
        text = paragraph.text.strip()
        if style.startswith("Heading") or style == "Title":
            if current in narrative_headings and len(seen) < 3:
                missing.append(current)
            current = text if text in narrative_headings else None
            seen = set()
            continue
        if current is None:
            continue
        for part in ("What we found.", "Why it matters.", "What to do."):
            if text.startswith(part.rstrip(".")):
                seen.add(part)
    if current in narrative_headings and len(seen) < 3:
        missing.append(current)
    return missing


SHAPE_COLUMN = ("Template", "Shape")


def _shape_column(header: List[str]) -> Optional[int]:
    """The one column allowed to name a template the way the tool does."""
    if header[:1] != [SHAPE_COLUMN[0]] or len(header) < 2:
        return None
    return 1 if header[1].startswith(SHAPE_COLUMN[1]) else None


def _client_texts(document) -> List[str]:
    """Everything a client reads, minus the appendix shape column.

    That column is the one place a template may be named the way the tool
    names it, because that is who it is for. Every other cell in every other
    table is client text, including the appendix tables.
    """
    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        header = [c.text.strip() for c in table.rows[0].cells]
        exempt = _shape_column(header)
        for index, row in enumerate(table.rows):
            for column, cell in enumerate(row.cells):
                if exempt is not None and index and column == exempt:
                    continue
                texts.append(cell.text)
    return texts


def _document_checks(document, path: str, expected_headings: List[str],
                     expected_images: int) -> List[str]:
    """The checks every format gets: order, images, dashes, shapes, size."""
    problems: List[str] = []

    headings = [p.text.strip() for p in document.paragraphs
                if p.style.name.startswith("Heading") or p.style.name == "Title"]
    position = 0
    for wanted in expected_headings:
        while position < len(headings) and headings[position] != wanted:
            position += 1
        if position >= len(headings):
            problems.append(f"heading missing or out of order: {wanted!r}")
        position += 1

    images = len(document.inline_shapes)
    if images != expected_images:
        problems.append(f"image count is {images}, expected {expected_images}")

    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)

    for text in texts:
        if _BAD_DASH.search(text):
            problems.append(f"dash used as punctuation: {text[:70]!r}")
            break
    for text in texts:
        lowered = text.lower()
        if any(p.lower() in lowered for p in PLACEHOLDERS):
            problems.append(f"placeholder text present: {text[:70]!r}")
            break

    # A template shape is right for the tool and wrong for a client. It is
    # allowed in one place only: the appendix column that exists to give the
    # site's developers the name the audit uses.
    for text in _client_texts(document):
        if "{" in text or "(depth" in text:
            problems.append(f"template shape in client text: {text[:70]!r}")
            break

    if os.path.getsize(path) > MAX_DOC_BYTES:
        problems.append(f"file is {os.path.getsize(path)} bytes, over the "
                        f"{MAX_DOC_BYTES} limit")
    return problems


def validate(path: str, expected_headings: List[str],
             expected_images: int) -> List[str]:
    """The long format: the shared checks plus its own narrative ones."""
    from docx import Document

    document = Document(path)
    problems = _document_checks(document, path, expected_headings,
                                expected_images)

    # Every heading must be followed by something before the next heading.
    body = [(p.text.strip(),
             p.style.name.startswith("Heading") or p.style.name == "Title")
            for p in document.paragraphs]
    for index, (text, is_heading) in enumerate(body):
        if not is_heading:
            continue
        following = body[index + 1:]
        has_body = any(t and not h for t, h in
                       following[:6]) if following else False
        if not has_body and index < len(body) - 1:
            # A heading may be followed by a table or an image instead.
            if not document.tables and not document.inline_shapes:
                problems.append(f"empty section body under {text!r}")

    # Truncation. The real failure mode is a list that stopped mid item, so
    # the check is on the list items themselves rather than on any paragraph
    # ending in a digit: "Trailing slash: 0." is data, not a cut off list.
    for paragraph in document.paragraphs:
        if paragraph.style.name != "List Number":
            continue
        item = paragraph.text.strip()
        if item and not _ENDS_PROPERLY.search(item):
            problems.append(f"action item ends mid sentence: {item[-60:]!r}")
            break

    # "What to do." is a label this code writes; it must be followed by at
    # least one action.
    paragraphs = list(document.paragraphs)
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() != "What to do.":
            continue
        following = paragraphs[index + 1:index + 2]
        if not following or following[0].style.name != "List Number":
            problems.append("a what to do heading has no actions under it")
            break

    # Every narrative section must carry all three parts.
    missing_parts = _sections_missing_parts(document)
    for heading in missing_parts:
        problems.append(f"section {heading!r} is missing a part")
    return problems


# --- validating the short report --------------------------------------------

# A count that states its whole: "3 of 842 pages parsed", "44 groups of 842
# pages parsed", "43.8 of 100 on a phone". The one allowed cell without two
# numbers is a setting that applies everywhere, which states its whole in
# words.
_STATES_A_WHOLE = re.compile(r"\d[\d,]*(?:\.\d+)?\s+(?:[a-z ]+\s+)?of\s+\d")
WHOLE_IN_WORDS = ("the whole site",)
_SENTENCE_END = re.compile(r"[.!?](?:\s|$)")


def _count_columns(table) -> List[int]:
    """The columns of this table that must state a whole, if any."""
    header = [c.text.strip() for c in table.rows[0].cells]
    return [index for index, name in enumerate(header)
            if name in ("Count", "Pages affected")]


def validate_short(path: str, expected_headings: List[str],
                   expected_images: int = SHORT_IMAGES) -> List[str]:
    """The shared checks, plus the ones that keep the short report short."""
    from docx import Document

    document = Document(path)
    problems = _document_checks(document, path, expected_headings,
                               expected_images)

    words = sum(len(p.text.split()) for p in document.paragraphs)
    if words > SHORT_MAX_WORDS:
        problems.append(f"{words} words outside tables, over the "
                        f"{SHORT_MAX_WORDS} limit")

    for table in document.tables:
        for column in _count_columns(table):
            for row in table.rows[1:]:
                cell = row.cells[column].text.strip()
                if row.cells[0].text.strip() == IN_PLACE:
                    continue
                if not cell:
                    problems.append(f"a count cell is empty in row "
                                    f"{row.cells[0].text[:40]!r}")
                    break
                if (not _STATES_A_WHOLE.search(cell)
                        and cell.lower() not in WHOLE_IN_WORDS):
                    problems.append(f"count states no whole: {cell[:70]!r}")
                    break

    for table in document.tables:
        header = [c.text.strip() for c in table.rows[0].cells]
        if header != FIXES_COLUMNS:
            continue
        for row in table.rows[1:]:
            if not row.cells[-1].text.strip():
                problems.append(f"fix has no action: "
                                f"{row.cells[1].text[:50]!r}")
                break

    # Prose is one lead sentence per section. The cover block is the one
    # place that runs longer, so the rule starts after it.
    started = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if paragraph.style.name.startswith("Heading") or \
                paragraph.style.name == "Title":
            if text == "Score":
                started = True
            continue
        if not started or not text:
            continue
        if len(_SENTENCE_END.findall(text)) > 2:
            problems.append(f"paragraph runs to prose: {text[:70]!r}")
            break
    return problems



def expected_headings_for(findings: Dict) -> List[str]:
    headings = ["Executive summary", "Terms used in this report",
                "Priority fixes"]
    headings += [title for key, title in SECTIONS if findings.get(key) is not None]
    if (findings.get("comparison") or {}).get("previous_run"):
        headings.append("Comparison with the previous audit")
    headings += ["Search performance", "Appendix"]
    return headings


# --- the command ------------------------------------------------------------

def _load_findings(run_dir: str) -> Dict:
    findings_path = os.path.join(run_dir, "findings.json")
    if not os.path.exists(findings_path):
        raise FileNotFoundError(
            f"no findings.json in {run_dir}; run the audit first")
    with open(findings_path, encoding="utf-8") as handle:
        return json.load(handle)


def generate_short(run_dir: str) -> Dict:
    """The client deliverable: no model, no cost, one validated file."""
    findings = _load_findings(run_dir)
    host = (findings.get("meta") or {}).get("host") or "site"

    charts = build_short_charts(findings)
    document = build_short_document(findings, charts, host)
    name = f"SEO_Audit_{host}_{date.today().isoformat()}.docx"
    path = os.path.join(run_dir, name)
    document.save(path)

    problems = validate_short(path, expected_headings_for_short(findings),
                              len(charts))
    if problems:
        raise ValidationError("; ".join(problems))
    return {"docx": path, "usage": None, "charts": len(charts),
            "format": "short", "usage_data": {"calls": 0, "model": "none",
                                              "input_tokens": 0,
                                              "output_tokens": 0,
                                              "max_calls": 0,
                                              "guard_events": []}}


def build_report(run_dir: str, fmt: str = "short", model: str = "gpt-5-mini",
                 use_ai: bool = True, session=None) -> Dict:
    """One entry point for both formats. Short unless asked otherwise."""
    if fmt == "short":
        return generate_short(run_dir)
    result = generate(run_dir, model=model, use_ai=use_ai, session=session)
    result["format"] = "long"
    return result


def generate(run_dir: str, model: str = "gpt-5-mini", use_ai: bool = True,
             session=None) -> Dict:
    """Build, write and validate the long report for one run folder."""
    findings = _load_findings(run_dir)

    host = (findings.get("meta") or {}).get("host") or "site"
    charts = build_charts(findings)
    parsed = ((findings.get("meta") or {}).get("coverage") or {}).get(
        "pages_parsed", 0)
    narrator = Narrator(model=model, use_ai=use_ai, session=session,
                        default_whole=parsed)

    document = build_document(findings, narrator, charts, host)
    name = f"SEO_Audit_{host}_{date.today().isoformat()}.docx"
    path = os.path.join(run_dir, name)
    document.save(path)

    usage_path = os.path.join(run_dir, "report_usage.json")
    with open(usage_path, "w", encoding="utf-8") as handle:
        json.dump(narrator.usage.as_dict(), handle, indent=2,
                  ensure_ascii=False)

    problems = validate(path, expected_headings_for(findings), len(charts))
    if problems:
        raise ValidationError("; ".join(problems))

    return {"docx": path, "usage": usage_path, "charts": len(charts),
            "usage_data": narrator.usage.as_dict()}


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="seo-audit-report",
        description="Build a Word report from a finished audit run folder.")
    parser.add_argument("--run", required=True,
                        help="Run folder containing findings.json")
    parser.add_argument("--model", default="gpt-5-mini",
                        help="Model for the narrative (default gpt-5-mini). "
                             "At most 14 calls per report.")
    parser.add_argument("--no-ai", action="store_true",
                        help="Long format only: write the narrative from "
                             "templates instead of calling the model.")
    parser.add_argument("--format", default="short", choices=("short", "long"),
                        dest="fmt",
                        help="short (default): the client deliverable, "
                             "tables and actions, no model calls. long: the "
                             "narrated report, for internal use.")
    return parser


def main(argv=None) -> int:
    from dotenv import load_dotenv

    args = build_parser().parse_args(argv)
    load_dotenv()

    try:
        result = build_report(args.run, fmt=args.fmt, model=args.model,
                              use_ai=not args.no_ai)
    except Exception as exc:
        print(f"{type(exc).__name__}: {exc}", file=sys.stderr)
        return 1

    usage = result["usage_data"]
    print(f"report   : {result['docx']}", file=sys.stderr)
    print(f"format   : {result['format']}", file=sys.stderr)
    print(f"charts   : {result['charts']}", file=sys.stderr)
    if result["format"] == "short":
        print("calls    : 0 (built from findings.json, no model)",
              file=sys.stderr)
        return 0
    print(f"usage    : {result['usage']}", file=sys.stderr)
    print(f"calls    : {usage['calls']} of {usage['max_calls']} "
          f"({usage['model']})", file=sys.stderr)
    print(f"tokens   : {usage['input_tokens']} in, "
          f"{usage['output_tokens']} out", file=sys.stderr)
    print(f"guarded  : {len(usage['guard_events'])} section(s) had a sentence "
          f"dropped", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
